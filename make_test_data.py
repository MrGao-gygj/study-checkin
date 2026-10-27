# -*- coding: utf-8 -*-
"""用样例数据验证模板里的公式计算是否正确（生成 test 副本，不动正式模板）。"""

import datetime

from openpyxl import load_workbook

SRC = r"E:\together study\打卡工具\学习打卡表.xlsx"
DST = r"E:\together study\打卡工具\test.xlsx"

wb = load_workbook(SRC)
ws = wb["打卡记录"]
members = wb["成员"]

today = datetime.date.today()
members["A2"] = "张三"
members["A3"] = "李四"

rows = [
    (today, "张三", "计算机", 2.0, 1.5, "21:10", "刷完概率论第 3 章", "已完成", "明天做真题", ""),
    (today, "李四", "数学", 3.0, 2.5, "22:05", "背 100 个单词 + 阅读", "部分完成", "补齐专业课作业", ""),
]
for i, row in enumerate(rows, start=2):
    for col, value in enumerate(row, start=1):
        if value != "":
            ws.cell(row=i, column=col, value=value)

wb.save(DST)
print("test copy saved:", DST)