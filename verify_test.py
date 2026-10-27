# -*- coding: utf-8 -*-
"""读取重算后的 test 副本，校验看板公式结果。"""

import datetime
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

wb = load_workbook(r"E:\together study\打卡工具\test_recalc.xlsx", data_only=True)
dash = wb["看板"]

today = datetime.datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
expected = {
    "B4": today,      # 今日日期
    "B5": 2,          # 应打卡人数
    "B6": 2,          # 今日打卡人数
    "B7": 1.0,        # 今日打卡率
    "B8": 5.0,        # 今日计划时长
    "B9": 4.0,        # 今日实际时长
    "B10": 5.0,       # 本周计划
    "B11": 4.0,       # 本周实际
    "B12": 2,         # 累计打卡次数
    "B13": 4.0,       # 累计实际时长
}

print("dashboard values:")
ok = True
for cell, want in expected.items():
    got = dash[cell].value
    if isinstance(want, datetime.datetime) and isinstance(got, datetime.datetime):
        passed = got.date() == want.date()
    else:
        passed = got == want
    if not passed:
        ok = False
    print(f"  {cell}: got={got!r} want={want!r} {'OK' if passed else 'MISMATCH'}")

print("\npersonal board:")
for r in range(6, 9):
    vals = [dash.cell(row=r, column=c).value for c in range(4, 8)]
    print("  ", vals)

print("\nverdict:", "PASS" if ok else "FAIL")