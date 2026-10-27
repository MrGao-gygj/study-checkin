# -*- coding: utf-8 -*-
"""生成学习打卡表模板（可直接导入腾讯文档 / 飞书 / WPS 使用）。"""

import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = r"E:\together study\打卡工具\学习打卡表.xlsx"

FONT = "微软雅黑"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT, bold=True, size=16, color="1F4E79")
LABEL_FONT = Font(name=FONT, bold=True, size=11)
BODY_FONT = Font(name=FONT, size=11)
SMALL_FONT = Font(name=FONT, size=10, color="666666")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

MAX_ROW = 1000

wb = Workbook()

# ---------------------------------------------------------------- 打卡记录
ws = wb.active
ws.title = "打卡记录"
headers = ["日期", "姓名", "专业/方向", "计划学习时长(小时)", "实际学习时长(小时)",
           "打卡时间", "今日学习任务", "完成情况", "明日计划", "备注"]
widths = [12, 12, 16, 16, 16, 12, 42, 12, 30, 20]
for col, (head, width) in enumerate(zip(headers, widths), start=1):
    cell = ws.cell(row=1, column=col, value=head)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = BORDER
    ws.column_dimensions[get_column_letter(col)].width = width
ws.row_dimensions[1].height = 30

for row in range(2, MAX_ROW + 1):
    ws.cell(row=row, column=1).number_format = "yyyy-mm-dd"
    ws.cell(row=row, column=4).number_format = "0.0"
    ws.cell(row=row, column=5).number_format = "0.0"
    ws.cell(row=row, column=7).alignment = WRAP
    ws.cell(row=row, column=9).alignment = WRAP
    ws.cell(row=row, column=1).alignment = CENTER
    ws.cell(row=row, column=2).alignment = CENTER
    ws.cell(row=row, column=6).alignment = CENTER
    ws.cell(row=row, column=8).alignment = CENTER

ws.freeze_panes = "A2"

# 完成情况下拉
dv_status = DataValidation(type="list", formula1='"已完成,部分完成,未完成"', allow_blank=True,
                           showErrorMessage=True, errorTitle="只能选择",
                           error="请选择：已完成 / 部分完成 / 未完成")
ws.add_data_validation(dv_status)
dv_status.add(f"H2:H{MAX_ROW}")

# 时长 0-24 校验
dv_hours = DataValidation(type="decimal", operator="between", formula1="0", formula2="24",
                          allow_blank=True, showErrorMessage=True, errorTitle="时长不合法",
                          error="请输入 0-24 之间的小时数")
ws.add_data_validation(dv_hours)
dv_hours.add(f"D2:E{MAX_ROW}")

# 今天这一行高亮
ws.conditional_formatting.add(
    f"A2:J{MAX_ROW}",
    FormulaRule(formula=[f"$A2=TODAY()"], fill=PatternFill("solid", fgColor="FFF2CC"), stopIfTrue=False),
)
# 完成情况配色
ws.conditional_formatting.add(
    f"H2:H{MAX_ROW}",
    CellIsRule(operator="equal", formula=['"已完成"'],
               fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=FONT, color="006100")),
)
ws.conditional_formatting.add(
    f"H2:H{MAX_ROW}",
    CellIsRule(operator="equal", formula=['"部分完成"'],
               fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(name=FONT, color="9C5700")),
)
ws.conditional_formatting.add(
    f"H2:H{MAX_ROW}",
    CellIsRule(operator="equal", formula=['"未完成"'],
               fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=FONT, color="9C0006")),
)

# 预填 5 行日期当天（方便直接打卡）——留空更干净，这里不预填数值

# ---------------------------------------------------------------- 看板
dash = wb.create_sheet("看板")
dash.sheet_view.showGridLines = False
dash.column_dimensions["A"].width = 26
dash.column_dimensions["B"].width = 14
dash.column_dimensions["C"].width = 34
dash.column_dimensions["D"].width = 14
dash.column_dimensions["E"].width = 12
dash.column_dimensions["F"].width = 16
dash.column_dimensions["G"].width = 14

dash.merge_cells("A1:G1")
t = dash["A1"]
t.value = "学习打卡看板"
t.font = TITLE_FONT
t.alignment = Alignment(horizontal="left", vertical="center")
dash.row_dimensions[1].height = 34

dash.merge_cells("A2:G2")
dash["A2"].value = "数据自动统计，不需要手动改；在「打卡记录」里新增一行即可。"
dash["A2"].font = SMALL_FONT

stats = [
    ("今日日期", "=TODAY()", "自动取今天的日期"),
    ("应打卡人数", "=COUNTA(成员!$A$2:$A$100)", "「成员」表里登记的名字数量"),
    ("今日打卡人数", "=COUNTIFS(打卡记录!$A$2:$A$1000,TODAY())", "今天已提交打卡的人数"),
    ("今日打卡率", "=IFERROR(B6/B5,0)", "今日打卡人数 ÷ 应打卡人数"),
    ("今日计划时长合计", "=SUMIFS(打卡记录!$D$2:$D$1000,打卡记录!$A$2:$A$1000,TODAY())", "单位：小时"),
    ("今日实际时长合计", "=SUMIFS(打卡记录!$E$2:$E$1000,打卡记录!$A$2:$A$1000,TODAY())", "单位：小时"),
    ("本周计划时长合计", '=SUMIFS(打卡记录!$D$2:$D$1000,打卡记录!$A$2:$A$1000,">="&TODAY()-WEEKDAY(TODAY(),2)+1,打卡记录!$A$2:$A$1000,"<="&TODAY())', "本周一到今天，单位：小时"),
    ("本周实际时长合计", '=SUMIFS(打卡记录!$E$2:$E$1000,打卡记录!$A$2:$A$1000,">="&TODAY()-WEEKDAY(TODAY(),2)+1,打卡记录!$A$2:$A$1000,"<="&TODAY())', "本周一到今天，单位：小时"),
    ("累计打卡次数", "=COUNTA(打卡记录!$B$2:$B$1000)", "所有历史打卡条数"),
    ("累计实际时长", "=SUM(打卡记录!$E$2:$E$1000)", "单位：小时"),
]

for i, (label, formula, note) in enumerate(stats):
    r = 4 + i
    dash.cell(row=r, column=1, value=label).font = LABEL_FONT
    cell = dash.cell(row=r, column=2, value=formula)
    cell.font = Font(name=FONT, bold=True, size=12, color="1F4E79")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    dash.cell(row=r, column=3, value=note).font = SMALL_FONT
    for c in range(1, 4):
        dash.cell(row=r, column=c).border = BORDER
    dash.row_dimensions[r].height = 24

dash["B4"].number_format = "yyyy-mm-dd"
dash["B7"].number_format = "0%"
for r in (8, 9, 10, 11, 13):
    dash.cell(row=r, column=2).number_format = "0.0"

# 个人榜
dash.merge_cells("D4:G4")
head = dash["D4"]
head.value = "个人本周榜（数据来自「成员」表）"
head.font = Font(name=FONT, bold=True, size=12, color="FFFFFF")
head.fill = HEADER_FILL
head.alignment = CENTER

board_headers = ["姓名", "本周打卡", "本周实际时长", "今日状态"]
for j, h in enumerate(board_headers):
    cell = dash.cell(row=5, column=4 + j, value=h)
    cell.font = LABEL_FONT
    cell.fill = PatternFill("solid", fgColor="DDEBF7")
    cell.alignment = CENTER
    cell.border = BORDER

for i in range(30):
    r = 6 + i
    member_row = 2 + i
    name_cell = dash.cell(row=r, column=4, value=f'=IF(成员!$A${member_row}="","",成员!$A${member_row})')
    name_cell.alignment = CENTER
    name_cell.border = BORDER
    week_start = "TODAY()-WEEKDAY(TODAY(),2)+1"
    dash.cell(row=r, column=5, value=(
        f'=IF($D{r}="","",COUNTIFS(打卡记录!$B$2:$B$1000,$D{r},'
        f'打卡记录!$A$2:$A$1000,">="&{week_start},打卡记录!$A$2:$A$1000,"<="&TODAY()))'
    )).alignment = CENTER
    dash.cell(row=r, column=6, value=(
        f'=IF($D{r}="","",SUMIFS(打卡记录!$E$2:$E$1000,打卡记录!$B$2:$B$1000,$D{r},'
        f'打卡记录!$A$2:$A$1000,">="&{week_start},打卡记录!$A$2:$A$1000,"<="&TODAY()))'
    )).alignment = CENTER
    dash.cell(row=r, column=7, value=(
        f'=IF($D{r}="","",IF(COUNTIFS(打卡记录!$B$2:$B$1000,$D{r},'
        f'打卡记录!$A$2:$A$1000,TODAY())>0,"✓ 已打卡","— 未打卡"))'
    )).alignment = CENTER
    for c in range(4, 8):
        dash.cell(row=r, column=c).border = BORDER
        dash.cell(row=r, column=c).font = BODY_FONT
    dash.cell(row=r, column=6).number_format = "0.0"

# ---------------------------------------------------------------- 成员
mem = wb.create_sheet("成员")
mem.sheet_view.showGridLines = False
mem_headers = ["姓名", "专业/方向", "目标/备注"]
mem_widths = [14, 22, 30]
for col, (h, w) in enumerate(zip(mem_headers, mem_widths), start=1):
    cell = mem.cell(row=1, column=col, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = BORDER
    mem.column_dimensions[get_column_letter(col)].width = w
mem.freeze_panes = "A2"
mem["A1"].comment = None

mem.merge_cells("E1:G1")
tip = mem["E1"]
tip.value = "把群成员名字填进 A 列，看板个人榜会自动出现；不用的人名留空即可。"
tip.font = SMALL_FONT
tip.alignment = Alignment(wrap_text=True, vertical="center")
mem.column_dimensions["E"].width = 14

# ---------------------------------------------------------------- 说明
ins = wb.create_sheet("说明")
ins.sheet_view.showGridLines = False
ins.column_dimensions["A"].width = 4
ins.column_dimensions["B"].width = 110

def line(row, text, font=None, height=None):
    cell = ins.cell(row=row, column=2, value=text)
    cell.font = font or BODY_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if height:
        ins.row_dimensions[row].height = height

ins.merge_cells("A1:B1")
ins["A1"].value = "使用说明（导入腾讯文档后，把本页删掉也不影响使用）"
ins["A1"].font = TITLE_FONT
ins.row_dimensions[1].height = 32

lines = [
    ("一、发到微信群（5 分钟）", True),
    ("1. 打开 腾讯文档（docs.qq.com 或微信小程序「腾讯文档」）→ 新建/导入 → 选择本文件「学习打卡表.xlsx」。", False),
    ("2. 右上角「分享」→ 权限选「所有人可编辑」（互联网上获得链接的人可编辑）→ 复制链接。", False),
    ("3. 把链接发到微信群。群成员点开即可在微信里直接编辑、查看，不需要装 App。", False),
    ("4. 手机端建议用「腾讯文档」小程序或 App，体验更顺；电脑端用网页即可。", False),
    ("", False),
    ("二、每天怎么打卡", True),
    ("1. 打开「打卡记录」表，在最后一行下面新增一行。", False),
    ("2. 填：日期（点日历选今天）、姓名、专业/方向、计划学习时长、实际学习时长、打卡时间、今日学习任务、完成情况（下拉选）、明日计划。", False),
    ("3. 保存即自动同步，「看板」里的今日人数、打卡率、时长合计会实时更新。", False),
    ("4. 想改自己那行，直接改；不要删别人的行。", False),
    ("", False),
    ("三、看板和成员表", True),
    ("· 「看板」自动统计：今日打卡人数/打卡率、今日与本周计划/实际时长、累计打卡次数和时长。", False),
    ("· 「成员」表 A 列填群成员姓名，看板右侧「个人本周榜」会自动生成每人本周打卡次数、时长和今日状态。", False),
    ("· 新成员加入：在「成员」表加一行名字即可。", False),
    ("", False),
    ("四、建议的群规则（可按需改）", True),
    ("· 每天 22:00 前完成打卡；每周允许补卡 1 次、可请假。", False),
    ("· 完成情况如实填写，不排名、不公开成绩，保护潜水成员。", False),
    ("· 周末用「看板」的周数据做一次 5 分钟复盘。", False),
    ("", False),
    ("五、常见问题", True),
    ("· 日期不是真日期（显示不对/统计为 0）：点日期单元格，从日历里重新选一次，或输入 2026-10-27 这样的格式。", False),
    ("· 想用飞书/WPS 也可以：同样导入 xlsx，飞书多维表格可把「打卡记录」转成表格视图+表单。", False),
    ("· 人数很多（>50）：建议每人每天只填一行，按日期筛选查看，避免表格过长。", False),
]

row = 3
for text, is_heading in lines:
    if not text:
        row += 1
        continue
    line(row, text, font=(Font(name=FONT, bold=True, size=12, color="1F4E79") if is_heading else BODY_FONT),
         height=(26 if is_heading else 22))
    row += 1

wb.save(OUT)
print("saved:", OUT)